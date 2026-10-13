#!/usr/bin/env python3
"""Genera i file Excel da compilare / caricare nel configuratore, nella cartella excel/.

  01_Matrice_Prezzi_Depureco.xlsx      prezzi macchine (solo colonna Prezzo modificabile) + accessori + commerciali
  02_Matrice_Optional_Depureco.xlsx    prezzi degli optional per macchina (matrice macchina x optional)
  03_Registro_Offerte_DEPURECO_2026.xlsx   registro offerte vuoto (alimenta la dashboard)
  04_Contatti_Clienti_Depureco.xlsx    anagrafica clienti per l'autocompletamento

Uso: python3 tools/build_excel_templates.py data/depureco_catalogo_AAAA-MM-GG.xlsx
"""
import subprocess, sys
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROSSO, GRIGIO, GIALLO = "AA1917", "171717", "FFF2B3"
RAD = Path(__file__).resolve().parent.parent
OUT = RAD / "excel"
GRUPPI = ["Filtrazione", "Struttura", "Costruzione", "Pulizia filtro", "Accessori", "Altro"]

def head(ws, riga, cols, colori=None):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=riga, column=j, value=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=(colori or {}).get(c, GRIGIO))
        cell.alignment = Alignment(vertical="center", wrap_text=True)

def titolo(ws, t, sub=None):
    ws["A1"] = t
    ws["A1"].font = Font(bold=True, size=14, color=ROSSO)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(italic=True, color="666666")

def istruzioni(wb, righe, in_fondo=False):
    """Foglio Istruzioni. in_fondo=True quando il configuratore legge il PRIMO foglio (file clienti)."""
    wi = wb.create_sheet("Istruzioni") if in_fondo else wb.create_sheet("Istruzioni", 0)
    for i, (t, b) in enumerate(righe, 1):
        c = wi.cell(row=i, column=1, value=t)
        c.font = Font(bold=b, size=14 if b else 11, color=ROSSO if b else "000000")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    wi.column_dimensions["A"].width = 120

def macchine(catalogo):
    ws = load_workbook(catalogo)["Macchine"]
    tutte = list(ws.iter_rows(values_only=True))
    i = next(i for i, r in enumerate(tutte[:15]) if r and "Codice" in r)
    return [dict(zip(tutte[i], r)) for r in tutte[i + 1:] if r and r[0]]

def matrice_optional(rows):
    sys.path.insert(0, str(RAD / "scraper"))
    import categorie
    rows = sorted(rows, key=lambda r: (categorie.ordine(r.get("Categoria") or ""), str(r.get("Sottocategoria") or ""), str(r.get("Nome") or "")))
    wb = Workbook()
    ws = wb.active
    ws.title = "Optional_Prezzi"
    # layout richiesto dal configuratore: riga "Gruppo", riga "Codice Macchina" (= codici optional), riga "Descrizione", poi una riga per macchina
    NCOL = 30
    ws.cell(row=1, column=1, value="Gruppo")
    ws.cell(row=2, column=1, value="Codice Macchina")
    ws.cell(row=3, column=1, value="Descrizione")
    for r in (1, 2, 3):
        ws.cell(row=r, column=1).font = Font(bold=True, color="FFFFFF")
        ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=GRIGIO)
    dv = DataValidation(type="list", formula1='"' + ",".join(GRUPPI) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    for j in range(2, NCOL + 2):
        for r in (1, 2, 3):
            c = ws.cell(row=r, column=j)
            c.fill = PatternFill("solid", fgColor=GIALLO)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        dv.add(f"{get_column_letter(j)}1")
        ws.column_dimensions[get_column_letter(j)].width = 16
        ws.cell(row=2, column=j).font = Font(bold=True)
    for n, r in enumerate(rows, 4):
        c = ws.cell(row=n, column=1, value=r["Codice"])
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="ECECEC")
        ws.cell(row=n, column=1).alignment = Alignment(vertical="center")
        for j in range(2, NCOL + 2):
            ws.cell(row=n, column=j).number_format = '#,##0.00 "€"'
    ws.column_dimensions["A"].width = 30
    ws.freeze_panes = "B4"
    ws.row_dimensions[3].height = 48

    wc = wb.create_sheet("Catalogo_Optional")
    head(wc, 1, ["Codice", "Gruppo", "Ordine", "Descrizione"])
    dv2 = DataValidation(type="list", formula1='"' + ",".join(GRUPPI) + '"', allow_blank=True)
    wc.add_data_validation(dv2)
    dv2.add("B2:B500")
    for j, w in enumerate((20, 18, 10, 80), 1):
        wc.column_dimensions[get_column_letter(j)].width = w
    wc.freeze_panes = "A2"

    wv = wb.create_sheet("Da_Verificare")
    head(wv, 1, ["Macchina", "Optional", "Nota"])
    for j, w in enumerate((24, 20, 60), 1):
        wv.column_dimensions[get_column_letter(j)].width = w

    istruzioni(wb, [
        ("DEPURECO — MATRICE OPTIONAL (macchina × optional)", True),
        ("", False),
        ("Come si compila (foglio Optional_Prezzi):", False),
        ("1. Riga 2 «Codice Macchina»: scrivere in ogni colonna il CODICE di un optional (una colonna per optional).", False),
        ("2. Riga 3 «Descrizione»: descrizione dell'optional. Riga 1 «Gruppo»: scegliere dal menu (Filtrazione, Struttura, Costruzione, Pulizia filtro, Accessori, Altro).", False),
        ("3. Nelle righe delle macchine (colonna A, già compilata) scrivere il PREZZO nella colonna dell'optional. Cella vuota = optional non disponibile per quella macchina.", False),
        ("4. Facoltativo: foglio Catalogo_Optional (codice, gruppo, ordine di visualizzazione, descrizione) e Da_Verificare (macchina + optional da controllare).", False),
        ("5. Nel configuratore: pulsante «Carica matrice optional» e scegliere questo file. Listino unico: nessuna distinzione L0/LE.", False),
        ("Le colonne gialle sono pronte per 30 optional: se ne servono di più, copiare l'ultima colonna a destra.", False),
    ])
    wb.save(OUT / "02_Matrice_Optional_Depureco.xlsx")

def registro():
    wb = Workbook()
    ws = wb.active
    ws.title = "Offerte"
    ws["A1"] = "DEPURECO — REGISTRO OFFERTE 2026"
    ws["A1"].font = Font(bold=True, size=14, color=ROSSO)
    head(ws, 3, ["Progressivo", "Numero_Offerta", "Data", "Cliente", "Referente", "Totale_EUR", "Macchina_Offerta",
                 "Optional_Aggiunti", "Note", "Stato", "Data_Esito", "Valore_Vinto"], {"Stato": ROSSO, "Valore_Vinto": ROSSO})
    for j, w in enumerate((12, 16, 12, 32, 26, 14, 50, 50, 30, 10, 12, 14), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A4"
    wr = wb.create_sheet("Righe")
    head(wr, 1, ["Numero_Offerta", "Pacchetto", "Codice", "Nome", "Categoria", "Qta", "Valore_EUR", "Venduta"])
    for j, w in enumerate((16, 11, 22, 32, 40, 8, 14, 10), 1):
        wr.column_dimensions[get_column_letter(j)].width = w
    istruzioni(wb, [
        ("DEPURECO — REGISTRO OFFERTE", True),
        ("", False),
        ("File gestito dal configuratore: non serve compilarlo a mano.", False),
        ("• Collegare la cartella che contiene questo file (pulsante «Collega cartella», Chrome o Edge): ogni «Salva offerta» aggiunge una riga.", False),
        ("• Il foglio «Offerte» ha una riga per offerta; «Righe» ha una riga per macchina offerta. La Dashboard legge entrambi.", False),
        ("• Stato: Aperta / Vinta / Persa (si imposta dalla Dashboard).", False),
        ("• Tenere UN solo registro condiviso tra i commerciali (disco di rete) per avere i numeri progressivi e la dashboard unificata.", False),
    ])
    wb.save(OUT / "03_Registro_Offerte_DEPURECO_2026.xlsx")

def contatti():
    wb = Workbook()
    ws = wb.active
    ws.title = "Clienti"
    head(ws, 1, ["Azienda", "Contatto", "Ruolo", "Email", "Telefono", "Indirizzo"])
    for j, w in enumerate((36, 26, 22, 34, 20, 44), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    istruzioni(wb, [
        ("DEPURECO — CONTATTI CLIENTI", True),
        ("", False),
        ("Una riga per cliente nel foglio Clienti (la colonna Azienda è obbligatoria; le altre sono facoltative).", False),
        ("Nel configuratore: «Carica clienti» e scegliere questo file. Nel passo Cliente il nome si completa da solo mentre si scrive.", False),
        ("Se la cartella è collegata, i clienti nuovi inseriti nelle offerte vengono aggiunti al file.", False),
        ("Il foglio «Clienti» deve restare il PRIMO foglio del file.", False),
    ], in_fondo=True)
    wb.save(OUT / "04_Contatti_Clienti_Depureco.xlsx")

def main():
    OUT.mkdir(exist_ok=True)
    catalogo = sys.argv[1]
    rows = macchine(catalogo)
    subprocess.run([sys.executable, str(RAD / "tools" / "build_price_matrix.py"), catalogo, "--out", str(OUT / "01_Matrice_Prezzi_Depureco.xlsx")], check=True)
    matrice_optional(rows)
    registro()
    contatti()
    print("excel/:", *sorted(p.name for p in OUT.glob("*.xlsx")), sep="\n  ")

if __name__ == "__main__":
    main()
