#!/usr/bin/env python3
"""Converte il catalogo Excel (foglio Macchine) in js/data/catalogo.js, caricato di default dal Generatore Offerte.
Uso: python3 tools/build_catalog_js.py data/depureco_catalogo_AAAA-MM-GG.xlsx
Stessa mappatura colonne->campi di js/app/xlsx-loader.js (handleXlsx)."""
import json, re, sys
from pathlib import Path
from openpyxl import load_workbook

CAMPI = {  # campo app -> colonna Excel
    "family": "Famiglia", "category": "Categoria", "name": "Nome", "desc": "Descrizione",
    "power": "Potenza", "voltage": "Tensione", "vacuum": "Depressione", "airflow": "Portata_Aria",
    "noise": "Rumorosita", "filter_type": "Tipo_Filtro", "filter_surface": "Sup_Filtrante",
    "capacity": "Capacita", "suction": "Bocca_Aspirazione", "dimensions": "Dimensioni", "weight": "Peso",
    "marking": "Marcatura_ATEX", "application": "Applicazione", "url": "URL", "image": "Immagine_URL",
}

def num(v):
    s = re.sub(r"[^\d.,-]", "", str(v or ""))
    s = re.sub(r"\.(?=\d{3}(\D|$))", "", s).replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0

def main(xlsx):
    ws = load_workbook(xlsx)["Macchine"]
    tutte = list(ws.iter_rows(values_only=True))
    i = next(i for i, r in enumerate(tutte[:15]) if r and "Codice" in r)   # salta titolo/sottotitolo
    righe = [dict(zip(tutte[i], r)) for r in tutte[i + 1:]]
    out = []
    for r in righe:
        if not r.get("Codice"):
            continue
        m = {"code": str(r["Codice"]).strip()}
        for k, col in CAMPI.items():
            m[k] = str(r.get(col) or "").strip()
        m["category"] = m["category"] or "ALTRO"
        m["name"] = m["name"] or m["code"]
        m["atex"] = bool(re.match(r"^(si|sì|yes|x|true|1|atex)$", str(r.get("ATEX") or ""), re.I))
        m["price"] = num(r.get("Prezzo"))
        out.append(m)
    dest = Path(__file__).resolve().parent.parent / "js" / "data" / "catalogo.js"
    dest.write_text("/* GENERATO da tools/build_catalog_js.py — non modificare a mano. Fonte: %s */\nconst CATALOGO_DEPURECO = %s;\n"
                    % (Path(xlsx).name, json.dumps(out, ensure_ascii=False)), encoding="utf8")
    print(f"{len(out)} modelli -> {dest}")

if __name__ == "__main__":
    main(sys.argv[1])
