#!/usr/bin/env python3
"""Genera la matrice prezzi Depureco: un Excel in cui si compila SOLO la colonna Prezzo.

Foglio Macchine: tutti i dati (codice, famiglia, nome, specifiche chiave) sono BLOCCATI; solo "Prezzo" (celle gialle) è modificabile.
Foglio Accessori: editabile (codici e prezzi degli accessori non sono sul sito).
Il Generatore Offerte la carica con "Carica listino": i prezzi si fondono col catalogo incorporato per Codice.
Uso: python3 tools/build_price_matrix.py data/depureco_catalogo_AAAA-MM-GG.xlsx [--out data/Matrice_Prezzi_Depureco.xlsx]
"""
import argparse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROSSO, GRIGIO, GIALLO = "AA1917", "171717", "FFF2B3"
COLONNE = ["Codice", "Famiglia", "Categoria", "Sottocategoria", "Protezione", "Nome", "ATEX", "Potenza", "Portata_Aria", "Depressione", "Capacita"]
LARG = {"Codice": 20, "Famiglia": 22, "Categoria": 38, "Sottocategoria": 22, "Protezione": 13, "Nome": 28, "ATEX": 8, "Potenza": 16, "Portata_Aria": 14,
        "Depressione": 14, "Capacita": 11, "Prezzo": 16}

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("catalogo")
    ap.add_argument("--out", default="data/Matrice_Prezzi_Depureco.xlsx")
    a = ap.parse_args()
    ws0 = load_workbook(a.catalogo)["Macchine"]
    tutte = list(ws0.iter_rows(values_only=True))
    i = next(i for i, r in enumerate(tutte[:15]) if r and "Codice" in r)
    righe = [dict(zip(tutte[i], r)) for r in tutte[i + 1:] if r and r[0]]
    import sys
    sys.path.insert(0, str(__import__("pathlib").Path(__file__).resolve().parent.parent / "scraper"))
    import categorie
    righe.sort(key=lambda r: (categorie.ordine(r.get("Categoria") or ""), str(r.get("Sottocategoria") or ""), str(r.get("Famiglia") or ""), str(r.get("Nome") or "")))

    wb = Workbook()
    wi = wb.active
    wi.title = "Istruzioni"
    testo = [
        ("DEPURECO — MATRICE PREZZI", True),
        ("", False),
        ("Cosa compilare: SOLO la colonna Prezzo (celle gialle) del foglio Macchine. Tutte le altre celle sono bloccate.", False),
        ("1. Per ogni macchina inserire il prezzo di listino in euro (numero, senza simbolo €, es. 4250 oppure 4250,50). Lasciare vuoto se il prezzo non è definito.", False),
        ("2. Foglio Accessori: aggiungere una riga per ogni accessorio/ricambio (Codice, Descrizione, Prezzo).", False),
        ("3. Salvare il file e, nel Generatore Offerte, cliccare «Carica listino» e scegliere questo file.", False),
        ("", False),
        ("Il listino è unico (non esistono L0/LE). I dati tecnici restano quelli del catalogo: il file porta solo i prezzi, abbinati per Codice.", False),
        ("Non cambiare i codici e non rinominare i fogli. Per sbloccare un foglio: Revisione → Rimuovi protezione foglio (nessuna password).", False),
    ]
    for r, (t, b) in enumerate(testo, 1):
        c = wi.cell(row=r, column=1, value=t)
        c.font = Font(bold=b, size=14 if b else 11, color=ROSSO if b else "000000")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    wi.column_dimensions["A"].width = 120

    ws = wb.create_sheet("Macchine")
    ws["A1"] = "DEPURECO — LISTINO MACCHINE: compilare solo la colonna PREZZO"
    ws["A1"].font = Font(bold=True, size=14, color=ROSSO)
    ws["A2"] = f"{len(righe)} modelli · listino unico in € · celle gialle = da compilare"
    ws["A2"].font = Font(italic=True, color="666666")
    cols = COLONNE + ["Prezzo"]
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=4, column=j, value=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=ROSSO if c == "Prezzo" else GRIGIO)
        cell.alignment = Alignment(vertical="center", horizontal="center" if c == "Prezzo" else "left")
        ws.column_dimensions[get_column_letter(j)].width = LARG[c]
    sottile = Side(style="thin", color="DDDDDD")
    for n, r in enumerate(righe, 5):
        for j, c in enumerate(COLONNE, 1):
            cell = ws.cell(row=n, column=j, value=r.get(c))
            cell.protection = Protection(locked=True)
            cell.font = Font(color="555555")
            cell.border = Border(bottom=sottile)
        p = ws.cell(row=n, column=len(cols))
        p.fill = PatternFill("solid", fgColor=GIALLO)
        p.number_format = '#,##0.00 "€"'
        p.protection = Protection(locked=False)
        p.border = Border(bottom=sottile)
    ult = 4 + len(righe)
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        errorTitle="Prezzo non valido", error="Inserire un numero maggiore o uguale a 0 (senza simbolo €).")
    ws.add_data_validation(dv)
    dv.add(f"{get_column_letter(len(cols))}5:{get_column_letter(len(cols))}{ult}")
    ws.freeze_panes = "G5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{ult}"
    ws.protection.sheet = True
    ws.protection.autoFilter = False   # filtri e ordinamento consentiti
    ws.protection.sort = False
    ws.protection.formatColumns = False

    wa = wb.create_sheet("Accessori")
    wa["A1"] = "DEPURECO — ACCESSORI E RICAMBI"
    wa["A1"].font = Font(bold=True, size=14, color=ROSSO)
    wa["A2"] = "Aggiungere una riga per accessorio: Codice, Descrizione, Prezzo (€)"
    wa["A2"].font = Font(italic=True, color="666666")
    for j, c in enumerate(["Codice", "Descrizione", "Prezzo"], 1):
        cell = wa.cell(row=4, column=j, value=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=ROSSO if c == "Prezzo" else GRIGIO)
    for j, w in enumerate((20, 80, 16), 1):
        wa.column_dimensions[get_column_letter(j)].width = w
    for n in range(5, 205):
        wa.cell(row=n, column=3).number_format = '#,##0.00 "€"'
        wa.cell(row=n, column=3).fill = PatternFill("solid", fgColor=GIALLO)
    wa.freeze_panes = "A5"
    wb.save(a.out)
    print(f"{len(righe)} modelli -> {a.out}")

if __name__ == "__main__":
    main()
