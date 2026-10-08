#!/usr/bin/env python3
"""Report di revisione categorie: per ogni famiglia, categoria del sito (come scritta) vs categoria assegnata.
Uso: python3 tools/report_categorie.py data/depureco_catalogo_AAAA-MM-GG.xlsx [--out data/Revisione_categorie.xlsx]"""
import argparse, collections, re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ap = argparse.ArgumentParser()
ap.add_argument("catalogo")
ap.add_argument("--out", default="data/Revisione_categorie.xlsx")
a = ap.parse_args()
ws0 = load_workbook(a.catalogo)["Macchine"]
tutte = list(ws0.iter_rows(values_only=True))
i = next(i for i, r in enumerate(tutte[:15]) if r and "Codice" in r)
righe = [dict(zip(tutte[i], r)) for r in tutte[i + 1:] if r and r[0]]
fam = collections.OrderedDict()
for r in righe:
    f = fam.setdefault(r["Famiglia"], {"sito": r.get("Categoria_Sito") or "(nessuna)", "cat": r["Categoria"], "sub": r["Sottocategoria"],
                                       "n": 0, "prot": collections.Counter(), "url": r["URL"]})
    f["n"] += 1
    f["prot"][r["Protezione"]] += 1

def primo_tag(sito):
    return re.split(r",\s*(?=Aspiratori|ASPIRATORI|Depolveratori|Impianti|Accessori)", sito)[0]

# categoria "equivalente" di ogni tag del sito = quella in cui finisce la maggioranza delle sue famiglie
maggioranza = {}
_c = collections.defaultdict(collections.Counter)
for f in fam.values():
    _c[primo_tag(f["sito"])][f["cat"]] += 1
for tag, cnt in _c.items():
    maggioranza[tag] = cnt.most_common(1)[0][0]

wb = Workbook()
ws = wb.active
ws.title = "Famiglie"
head = ["Famiglia", "N_modelli", "Categoria_nuova", "Sottocategoria", "Protezioni", "Categoria_sito (come scritta)", "Esito", "URL"]
ws.append(["REVISIONE CATEGORIE — famiglie Depureco (verifica sezione corretta)"])
ws["A1"].font = Font(bold=True, size=14, color="AA1917")
ws.append([])
ws.append(head)
for j in range(1, len(head) + 1):
    ws.cell(row=3, column=j).font = Font(bold=True, color="FFFFFF")
    ws.cell(row=3, column=j).fill = PatternFill("solid", fgColor="171717")
COLORE = {"invariata": "C6EFCE", "spostata": "F8CBAD", "assegnata": "D9E7FF"}
for nome, f in fam.items():
    tag = primo_tag(f["sito"])
    if f["sito"] == "(nessuna)":
        esito, col = "assegnata (il sito non aveva categoria)", "assegnata"
    elif f["cat"] == maggioranza[tag] and "Atex ed ACD" not in tag:
        esito, col = "stessa categoria (rinominata)" if len(primo_tag(f["sito"])) else "invariata", "invariata"
    else:
        esito, col = f"spostata da «{tag}»", "spostata"
    ws.append([nome, f["n"], f["cat"], f["sub"], ", ".join(f"{k} {v}" for k, v in f["prot"].items()), f["sito"], esito, f["url"]])
    ws.cell(row=ws.max_row, column=7).fill = PatternFill("solid", fgColor=COLORE[col])
for j, w in enumerate((30, 10, 44, 26, 30, 70, 36, 60), 1):
    ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = "B4"
ws.auto_filter.ref = f"A3:H{ws.max_row}"
wb.save(a.out)
c = collections.Counter(ws.cell(row=r, column=7).value.split(" da ")[0] for r in range(4, ws.max_row + 1))
print(f"{len(fam)} famiglie -> {a.out}\n{dict(c)}")
