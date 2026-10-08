#!/usr/bin/env python3
"""Aggiunge il foglio 'Confronto' (Dupuy vs Depureco) al file catalogo Depureco.

Per ogni modello Dupuy cerca i 3 modelli Depureco piu' vicini per potenza, portata d'aria,
depressione e capacita' (stesso stato ATEX obbligatorio). Depressione Dupuy convertita mm H2O -> mbar.
Uso: python3 confronta_dupuy_depureco.py LISTINO_DUPUY.xlsx CATALOGO_DEPURECO.xlsx [--out FILE]
"""
import argparse, math, re
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

MMH2O_TO_MBAR = 0.0980665  # 1 mm H2O = 0.0980665 mbar
PESI = {"pot": 1.0, "port": 1.0, "vuoto": 1.0, "cap": 0.5}


def num(s):
    """Primo numero di una stringa ('2.400 mm H2O' -> 2400, '2 x 1,1 bypass kW' -> 2.2)."""
    if not s:
        return None
    s = str(s)
    m = re.match(r"\s*(\d+)\s*x\s*(\d+(?:,\d+)?)", s)
    if m:
        return int(m.group(1)) * float(m.group(2).replace(",", "."))
    m = re.search(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?|\d+(?:,\d+)?", s)
    if not m:
        return None
    t = m.group(0)
    t = t.replace(".", "").replace(",", ".") if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", t) else t.replace(",", ".")
    return float(t)


def metriche(pot, port, vuoto, cap, vuoto_mmh2o=False):
    v = num(vuoto)
    if v is not None and vuoto_mmh2o:
        v *= MMH2O_TO_MBAR
    return {"pot": num(pot), "port": num(port), "vuoto": v, "cap": num(cap)}


def distanza(a, b):
    tot, pesi, n = 0.0, 0.0, 0
    for k, w in PESI.items():
        if a[k] and b[k]:
            tot += w * abs(math.log(b[k] / a[k]))
            pesi += w
            if k != "cap":
                n += 1
    return (tot / pesi if pesi else None), n


def pct(a, b):
    return (b - a) / a if a and b else None


def giudizio(deltas):
    d = [abs(x) for x in deltas if x is not None]
    if len(d) < 2:
        return "Dati insufficienti"
    m = max(d)
    return "Equivalente" if m <= 0.15 else "Simile" if m <= 0.35 else "Approssimato"


def leggi(ws, hdr_row):
    h = [c.value for c in ws[hdr_row]]
    return [dict(zip(h, r)) for r in ws.iter_rows(min_row=hdr_row + 1, values_only=True) if any(r)]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("dupuy")
    ap.add_argument("depureco")
    ap.add_argument("--out", default=None)
    a = ap.parse_args()
    dup = leggi(load_workbook(a.dupuy)["Macchine"], 4)
    wb = load_workbook(a.depureco)
    dep = leggi(wb["Macchine"], 4)
    for d in dep:
        d["_m"] = metriche(d["Potenza"], d["Portata_Aria"], d["Depressione"], d["Capacita"])
        d["_atex"] = bool(d["ATEX"])

    cols = ["Dupuy_Codice", "Dupuy_Famiglia", "Dupuy_Nome", "ATEX", "Prezzo_L0", "Prezzo_LE",
            "D_Potenza_kW", "D_Portata_m3h", "D_Depress_mbar", "D_Capacita_lt",
            "Giudizio", "Depureco_Nome", "Depureco_Famiglia", "P_Potenza_kW", "P_Portata_m3h",
            "P_Depress_mbar", "P_Capacita_lt", "Δ_Potenza_%", "Δ_Portata_%", "Δ_Depress_%",
            "Alternativa_2", "Alternativa_3", "Depureco_URL"]
    if "Confronto" in wb.sheetnames:
        del wb["Confronto"]
    ws = wb.create_sheet("Confronto", 0)
    ws["A1"] = "CONFRONTO DUPUY vs DEPURECO — modello Depureco più vicino per prestazioni"
    ws["A1"].font = Font(bold=True, size=14, color="152239")
    ws["A2"] = ("Match per potenza, portata d'aria, depressione (mm H2O Dupuy → mbar), capacità; stesso stato ATEX. "
                "Giudizio: Equivalente = tutti i Δ ≤15% · Simile ≤35% · Approssimato >35%. "
                "Δ% = (Depureco − Dupuy) / Dupuy. Prezzi Depureco non pubblici.")
    ws["A2"].font = Font(italic=True, color="666666")
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=4, column=i, value=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="152239")
    colori = {"Equivalente": "C6EFCE", "Simile": "FFEB9C", "Approssimato": "F8CBAD", "Dati insufficienti": "D9D9D9"}
    for r in dup:
        atex = bool(r["ATEX"])
        m = metriche(r["Potenza"], r["Portata_Aria"], r["Depressione"], r["Capacita"], vuoto_mmh2o=True)
        cand = []
        for d in dep:
            if d["_atex"] != atex:
                continue
            dist, n = distanza(m, d["_m"])
            if dist is not None and n >= 2:
                cand.append((dist, d))
        cand.sort(key=lambda x: x[0])
        riga = [r["Codice"], r["Famiglia"], r["Nome"], "SI" if atex else "", r["Prezzo_L0"], r["Prezzo_LE"],
                m["pot"], m["port"], round(m["vuoto"]) if m["vuoto"] else None, m["cap"]]
        if cand:
            b = cand[0][1]["_m"]
            dl = [pct(m["pot"], b["pot"]), pct(m["port"], b["port"]), pct(m["vuoto"], b["vuoto"])]
            bd = cand[0][1]
            g = giudizio(dl)
            riga += [g, bd["Nome"], bd["Famiglia"], b["pot"], b["port"], round(b["vuoto"]) if b["vuoto"] else None,
                     b["cap"], *dl]
            alt = [f'{c[1]["Nome"]} ({c[1]["Potenza"] or "?"} · {c[1]["Portata_Aria"] or "?"})' for c in cand[1:3]]
            riga += alt + [""] * (2 - len(alt)) + [bd["URL"]]
        else:
            g = "Dati insufficienti"
            riga += [g] + [""] * (len(cols) - len(riga) - 1)
        ws.append(riga)
        ws.cell(row=ws.max_row, column=11).fill = PatternFill("solid", fgColor=colori[g])
        for ci in (18, 19, 20):
            ws.cell(row=ws.max_row, column=ci).number_format = "+0%;-0%;0%"
    ws.freeze_panes = "D5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{ws.max_row}"
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = 34 if c.startswith("Alternativa") or c == "Depureco_URL" else 16
    out = a.out or a.depureco.replace(".xlsx", "_vs_dupuy.xlsx")
    wb.save(out)
    tot = ws.max_row - 4
    from collections import Counter
    cnt = Counter(ws.cell(row=i, column=11).value for i in range(5, ws.max_row + 1))
    print(f"{tot} modelli Dupuy confrontati -> {out}\n{dict(cnt)}")


if __name__ == "__main__":
    main()
