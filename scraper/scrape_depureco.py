#!/usr/bin/env python3
"""Scraper catalogo Depureco -> Excel con la stessa struttura del listino Dupuy (foglio Macchine).

Fonte: sitemap Yoast (prodotto-sitemap.xml) + JSON-LD schema.org Product/ProductModel di ogni scheda.
Solo libreria standard + openpyxl. Rispetta robots.txt (nessun path vietato usato) con pausa tra le richieste.
Uso: python3 scrape_depureco.py [--out FILE.xlsx] [--limit N] [--delay 0.6] [--images DIR] [--img-size 500]
"""
import argparse, html as htmllib, json, os, re, shutil, sys, time, urllib.request
from datetime import date
from xml.etree import ElementTree as ET
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE = "https://www.depureco.com"
SITEMAP = BASE + "/prodotto-sitemap.xml"
IT_PREFIX = BASE + "/aspiratori-industriali/"
UA = "Mozilla/5.0 (compatible; DupuyCompetitiveIntel/1.0; +mailto:r.armentaro@dupuy.it)"

COLONNE = ["Codice", "Famiglia", "Categoria", "Nome", "ATEX", "Prezzo", "Potenza", "Tensione",
           "Depressione", "Portata_Aria", "Rumorosita", "Tipo_Filtro", "Sup_Filtrante", "Capacita",
           "Bocca_Aspirazione", "Dimensioni", "Peso", "Altezza", "Marcatura_ATEX", "Zone_ATEX",
           "Applicazione", "Altre_Categorie", "Altre_Specifiche", "Descrizione", "URL", "Immagine_URL", "Immagine_File"]

# colonna -> nomi proprieta' (minuscoli) sul sito
ALIAS = {
    "Potenza": ["potenza", "power"],
    "Tensione": ["voltaggio motore"],
    "Depressione": ["vuoto in continuo"],
    "Portata_Aria": ["massima portata d'aria"],
    "Rumorosita": ["livello di rumorosità"],
    "Tipo_Filtro": ["tipologia filtro primario"],
    "Sup_Filtrante": ["superficie filtrante"],
    "Capacita": ["capacità"],
    "Bocca_Aspirazione": ["bocca aspirante"],
    "Dimensioni": ["dimensioni"],
    "Peso": ["peso"],
    "Altezza": ["altezza"],
    "Marcatura_ATEX": ["marcatura atex"],
}
# gia' mappate o ridondanti: non finiscono in Altre_Specifiche
SKIP_EXTRA = {a for v in ALIAS.values() for a in v} | {"elenco puntato", "tipologie", "frequenza"}


def get(url):
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    for tentativo in range(3):
        try:
            with urllib.request.urlopen(req, timeout=30) as r:
                return r.read().decode("utf-8", "ignore")
        except Exception as e:
            if tentativo == 2:
                raise
            time.sleep(2 * (tentativo + 1))


def urls_prodotti():
    root = ET.fromstring(get(SITEMAP))
    ns = {"s": "http://www.sitemaps.org/schemas/sitemap/0.9"}
    return sorted({e.text.strip() for e in root.findall(".//s:loc", ns) if e.text.strip().startswith(IT_PREFIX)})


def jsonld_product(html):
    for m in re.finditer(r'<script type="application/ld\+json"[^>]*>(.*?)</script>', html, re.S):
        try:
            d = json.loads(m.group(1))
        except ValueError:
            continue
        stack = [d]
        while stack:
            o = stack.pop()
            if isinstance(o, dict):
                if o.get("@type") == "Product" and o.get("model") is not None or o.get("@type") == "Product":
                    return o
                stack.extend(o.values())
            elif isinstance(o, list):
                stack.extend(o)
    return None


def props(lst):
    return {p["name"].strip().lower(): str(p["value"]).strip() for p in lst or [] if "name" in p and "value" in p}


def pulisci(t):
    return re.sub(r"\s+", " ", t or "").strip()


def righe_da_pagina(url, html):
    p = jsonld_product(html)
    if not p:
        return []
    pp = props(p.get("additionalProperty"))
    modelli = p.get("model") or [None]
    righe = []
    for m in modelli:
        mp = props(m.get("additionalProperty")) if m else {}
        allp = {**pp, **mp}  # le specifiche del modello prevalgono
        nome = (m or {}).get("name") or p["name"]
        r = {c: "" for c in COLONNE}
        r["Codice"] = re.sub(r"[^A-Z0-9]", "", nome.upper())
        r["Famiglia"] = p["name"]
        cats = [c.strip() for c in re.split(r",\s*(?=Aspiratori|ASPIRATORI|Depolveratori|Impianti|Accessori)", pp.get("tipologie", "")) if c.strip()]
        r["Categoria"] = cats[0] if cats else ""
        r["Altre_Categorie"] = ", ".join(cats[1:])
        img = p.get("image")
        img = img[0] if isinstance(img, list) and img else img
        r["Immagine_URL"] = (img.get("url") if isinstance(img, dict) else img) or ""
        r["Nome"] = nome
        zone = pp.get("tipi zona", "")
        cert = pp.get("certificati", "")
        r["ATEX"] = "ATEX" if ("atex" in (cert + zone + allp.get("marcatura atex", "")).lower()) else ""
        for col, nomi in ALIAS.items():
            for n in nomi:
                if allp.get(n):
                    r[col] = allp[n]
                    break
        r["Zone_ATEX"] = zone if r["ATEX"] else ""
        r["Applicazione"] = pp.get("applicazione", "")
        extra = [f"{k}: {v}" for k, v in allp.items() if k not in SKIP_EXTRA
                 and k not in ("applicazione", "tipi zona", "certificati", "alimentazioni", "tipo prodotto", "utilizzi")]
        r["Altre_Specifiche"] = " | ".join(extra)
        r["Descrizione"] = pulisci(p.get("description"))
        r["URL"] = p.get("url") or url
        righe.append({k: htmllib.unescape(v) if isinstance(v, str) else v for k, v in r.items()})
    return righe


def scarica_immagini(righe, cartella, size, delay):
    """Scarica l'immagine di ogni modello e la salva come <Codice>.<ext> in `cartella`.
    Le immagini condivise da piu' modelli (stessa famiglia) vengono scaricate una volta e copiate col nome di ogni modello.
    Imposta r["Immagine_File"]. Ritorna (scaricate, copiate, errori)."""
    os.makedirs(cartella, exist_ok=True)
    cache, ok, copie, errori = {}, 0, 0, []
    for r in righe:
        url = r.get("Immagine_URL")
        if not url:
            errori.append((r["Codice"], "nessuna immagine sul sito"))
            continue
        url = re.sub(r"fit=\d+%2C\d+", f"fit={size}%2C{size}", url)  # versione ridimensionata servita dal CDN del sito
        try:
            if url not in cache:
                req = urllib.request.Request(url, headers={"User-Agent": UA})
                with urllib.request.urlopen(req, timeout=30) as resp:
                    tipo = resp.headers.get("Content-Type", "")
                    dati = resp.read()
                if not tipo.startswith("image/") or len(dati) < 500:
                    raise ValueError(f"risposta non valida ({tipo}, {len(dati)} byte)")
                ext = {"image/png": "png", "image/jpeg": "jpg", "image/webp": "webp", "image/gif": "gif"}.get(tipo.split(";")[0], "png")
                nome = f"{r['Codice']}.{ext}"
                with open(os.path.join(cartella, nome), "wb") as f:
                    f.write(dati)
                cache[url] = (nome, dati, ext)
                ok += 1
                time.sleep(delay)
            else:
                _, dati, ext = cache[url]
                nome = f"{r['Codice']}.{ext}"
                with open(os.path.join(cartella, nome), "wb") as f:
                    f.write(dati)
                copie += 1
            r["Immagine_File"] = nome
        except Exception as e:
            errori.append((r["Codice"], repr(e)))
    return ok, copie, errori


def salva(righe, out):
    """Cartella di lavoro compatibile col Generatore Offerte: fogli Macchine, Accessori, Commerciali, Istruzioni."""
    ROSSO, GRIGIO = "AA1917", "171717"  # brand Depureco
    wb = Workbook()
    ws = wb.active
    ws.title = "Macchine"
    ws["A1"] = "DEPURECO — CATALOGO MACCHINE (scraping www.depureco.com)"
    ws["A1"].font = Font(bold=True, size=14, color=ROSSO)
    ws["A2"] = (f"Estratto il {date.today().isoformat()} · {len(righe)} modelli · prezzi non pubblicati dal sito: "
                "compilare la colonna Prezzo (listino unico) · struttura del listino Dupuy senza L0/LE")
    ws["A2"].font = Font(italic=True, color="666666")
    def intestazione(sheet, cols, riga):
        for i, c in enumerate(cols, 1):
            cell = sheet.cell(row=riga, column=i, value=c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=GRIGIO if i % 2 else ROSSO)
            cell.alignment = Alignment(vertical="center")
    intestazione(ws, COLONNE, 4)
    for r in righe:
        ws.append([r[c] for c in COLONNE])
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(COLONNE))}{4 + len(righe)}"
    larg = {"Altre_Specifiche": 50, "Descrizione": 60, "URL": 55, "Immagine_URL": 55, "Immagine_File": 28, "Applicazione": 30,
            "Categoria": 28, "Altre_Categorie": 28}
    for i, c in enumerate(COLONNE, 1):
        ws.column_dimensions[get_column_letter(i)].width = larg.get(c, 18)

    wa = wb.create_sheet("Accessori")
    wa["A1"] = "DEPURECO — LISTINO ACCESSORI"
    wa["A1"].font = Font(bold=True, size=14, color=ROSSO)
    wa["A2"] = "Accessori, kit e ricambi · da compilare dal listino interno (il sito non pubblica codici/prezzi)"
    intestazione(wa, ["Codice", "Descrizione", "Prezzo"], 4)
    for col, w in zip("ABC", (18, 70, 14)):
        wa.column_dimensions[col].width = w

    wc = wb.create_sheet("Commerciali")
    wc["A1"] = "DEPURECO — REFERENTI COMMERCIALI"
    wc["A1"].font = Font(bold=True, size=14, color=ROSSO)
    intestazione(wc, ["Nome", "Email", "Telefono"], 3)
    wc.append(["Ufficio Commerciale Depureco", "depureco@depureco.com", "+39 011 9859117"])  # contatti pubblici del sito
    for col, w in zip("ABC", (32, 34, 22)):
        wc.column_dimensions[col].width = w

    wi = wb.create_sheet("Istruzioni")
    for i, riga in enumerate([
        "DEPURECO — FILE SORGENTE CATALOGO PER IL GENERATORE OFFERTE",
        "",
        "Generato da scraper/scrape_depureco.py (fonte: www.depureco.com, schede prodotto pubbliche).",
        "1. Compilare la colonna Prezzo (listino unico) nel foglio Macchine.",
        "2. Compilare il foglio Accessori (Codice, Descrizione, Prezzo) e aggiornare Commerciali.",
        "3. Nel Generatore Offerte: pulsante 'Carica listino' e scegliere questo file.",
        "Le colonne Immagine_URL, URL, Altre_Categorie e Altre_Specifiche sono informative.",
    ], 1):
        wi.cell(row=i, column=1, value=riga)
    wi["A1"].font = Font(bold=True, size=13, color=ROSSO)
    wi.column_dimensions["A"].width = 110
    wb.save(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=f"depureco_catalogo_{date.today().isoformat()}.xlsx")
    ap.add_argument("--limit", type=int, default=0)
    ap.add_argument("--delay", type=float, default=0.6)
    ap.add_argument("--images", default="", help="cartella dove salvare le immagini (<Codice>.png); vuoto = non scaricare")
    ap.add_argument("--img-size", type=int, default=500, help="lato massimo in px delle immagini (default 500)")
    a = ap.parse_args()
    urls = urls_prodotti()
    if a.limit:
        urls = urls[: a.limit]
    print(f"{len(urls)} schede prodotto trovate in sitemap", file=sys.stderr)
    righe, errori, viste = [], [], set()
    for i, u in enumerate(urls, 1):
        try:
            rr = righe_da_pagina(u, get(u))
            if not rr:
                errori.append((u, "nessun JSON-LD Product"))
            for r in rr:
                k = (r["Famiglia"], r["Nome"])
                if k not in viste:
                    viste.add(k)
                    righe.append(r)
            print(f"[{i}/{len(urls)}] {len(rr)} modelli · {u.rsplit('/', 2)[-2]}", file=sys.stderr)
        except Exception as e:
            errori.append((u, repr(e)))
        time.sleep(a.delay)
    if a.images:
        ok, copie, err_img = scarica_immagini(righe, a.images, a.img_size, a.delay / 2)
        print(f"Immagini: {ok} scaricate + {copie} copie per modelli della stessa famiglia -> {a.images}/", file=sys.stderr)
        errori += [(c, "immagine: " + e) for c, e in err_img]
    salva(righe, a.out)
    print(f"\nOK: {len(righe)} modelli -> {a.out}")
    for u, e in errori:
        print(f"ERRORE {u}: {e}")


if __name__ == "__main__":
    main()
